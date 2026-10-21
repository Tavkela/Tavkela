"""Casco: სურათების ჩამოწერა casco-helmets.com-დან Excel-ის არტიკულების მიხედვით.

გამოყენება:
    python3 casco_images.py Casco.xlsx --out OUT_DIR [--size 1500] [--odoo https://ODOO_DOMAIN]

შედეგი OUT_DIR-ში:
    images/<კოდი>.jpg, <კოდი>_A.jpg, ...   (პირველი: მთავარი სურათი თეთრ ფონზე)
    Casco - შედეგი.xlsx                       (მწვანე/ყვითელი/წითელი/ლურჯი მონიშვნით)
"""
import argparse
import io
import json
import re
import shutil
import sys
import time
from pathlib import Path
from urllib.parse import quote

import openpyxl
import requests
from openpyxl.styles import PatternFill
from PIL import Image

SITE = "https://casco-helmets.com"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/141.0 Safari/537.36"}

FILLS = {
    "green": PatternFill("solid", fgColor="C6EFCE"),
    "yellow": PatternFill("solid", fgColor="FFEB9C"),
    "red": PatternFill("solid", fgColor="FFC7CE"),
    "blue": PatternFill("solid", fgColor="BDD7EE"),
}

ARTICLE_RE = re.compile(r"(\d{2})\.(\d{4})")


def base_article(code):
    """'07.2015.M.ML' -> '07.2015'"""
    m = ARTICLE_RE.search(str(code or "").replace(",", "."))
    return f"{m.group(1)}.{m.group(2)}" if m else None


def suffix(i):
    """0 -> '', 1 -> '_A', 26 -> '_Z', 27 -> '_AA'"""
    if i == 0:
        return ""
    s = ""
    while i > 0:
        i -= 1
        s = chr(65 + i % 26) + s
        i //= 26
    return "_" + s


def get(url, session, **kw):
    for attempt in range(5):
        r = session.get(url, headers=UA, timeout=60, **kw)
        if r.status_code in (429, 502, 503):
            time.sleep(2 ** attempt)
            continue
        r.raise_for_status()
        return r
    r.raise_for_status()
    return r


# ---------- Casco-ს კატალოგი (Shopify) ----------

def load_catalog(session, cache):
    if cache.exists():
        return json.loads(cache.read_text(encoding="utf-8"))
    products, page = [], 1
    while True:
        r = get(f"{SITE}/products.json?limit=250&page={page}", session)
        batch = r.json().get("products", [])
        if not batch:
            break
        products += batch
        print(f"  კატალოგი: გვერდი {page}, სულ {len(products)} პროდუქტი")
        page += 1
        time.sleep(0.5)
    cache.write_text(json.dumps(products, ensure_ascii=False), encoding="utf-8")
    return products


def image_name(src):
    return src.split("?")[0].rsplit("/", 1)[-1]


def build_index(products):
    """არტიკული -> [(პროდუქტი, წყარო)]; წყარო: sku > სურათის სახელი > handle."""
    index = {}

    def add(art, p, how):
        lst = index.setdefault(art, [])
        if all(q["id"] != p["id"] for q, _ in lst):
            lst.append((p, how))

    for p in products:
        for v in p.get("variants", []):
            a = base_article(v.get("sku"))
            if a:
                add(a, p, "sku")
    for p in products:
        for img in p.get("images", []):
            m = re.match(r"(\d{2}\.\d{4})_", image_name(img["src"]))
            if m:
                add(m.group(1), p, "image")
    for p in products:
        m = re.search(r"(\d{2})-(\d{4})(?:-|$)", p.get("handle", ""))
        if m:
            add(f"{m.group(1)}.{m.group(2)}", p, "handle")
    for lst in index.values():
        lst.sort(key=lambda t: ["sku", "image", "handle"].index(t[1]))
    return index


def select_images(product, art):
    """პროდუქტის სურათები ამ არტიკულისთვის, თანმიმდევრობით (position)."""
    imgs = sorted(product.get("images", []), key=lambda i: i.get("position", 0))
    arts = {base_article(v.get("sku")) for v in product.get("variants", [])} - {None}
    if len(arts) <= 1:
        chosen = imgs
    else:
        # ერთ პროდუქტში რამდენიმე ფერია: ვიღებთ მხოლოდ ამ არტიკულის სურათებს
        vids = {v["id"] for v in product["variants"] if base_article(v.get("sku")) == art}
        chosen = [i for i in imgs
                  if image_name(i["src"]).startswith(art + "_")
                  or vids & set(i.get("variant_ids") or [])]
        chosen = chosen or imgs
    seen, out = set(), []
    for i in chosen:
        if i["src"] not in seen:
            seen.add(i["src"])
            out.append(i["src"])
    return out


def tokens(s):
    s = re.sub(r"[^A-Z0-9]+", " ", str(s).upper())
    out = []
    for t in s.split():
        if t in {"ML", "MAGNET", "LINK", "CASCO"}:
            continue
        out.append(re.sub(r"(.)\1+", r"\1", t))  # DUELL -> DUEL, MISTRALL -> MISTRAL
    return out


def name_candidates(model, products, limit=3):
    want = tokens(model)
    if not want:
        return []
    found = []
    for p in products:
        have = set(tokens(p.get("title", "")) + tokens(p.get("handle", "")))
        if all(t in have for t in want):
            found.append(p)
    return found[:limit]


# ---------- Odoo-ს საჯარო საიტი ----------

def odoo_has(session, odoo, code):
    """True, თუ Odoo-ს საიტზე ამ კოდით პროდუქტი იძებნება."""
    r = get(f"{odoo}/shop?search={quote(str(code))}", session)
    links = set(re.findall(r'href="(/shop/(?!cart|checkout|category|page|wishlist)[^"?#]+-\d+)"', r.text))
    return len(links) > 0


# ---------- სურათები ----------

def fetch_jpeg(session, src, size):
    url = "https:" + src if src.startswith("//") else src
    if size:
        url += ("&" if "?" in url else "?") + f"width={size}"
    data = get(url, session).content
    im = Image.open(io.BytesIO(data))
    im.load()
    if size and max(im.size) > size:
        im.thumbnail((size, size), Image.LANCZOS)
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.split()[-1])
        im = bg
    elif im.mode != "RGB":
        im = im.convert("RGB")
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=92, optimize=True)
    return buf.getvalue()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--out", required=True)
    ap.add_argument("--size", type=int, default=1500, help="0 = ორიგინალი ზომა")
    ap.add_argument("--odoo", default=None)
    args = ap.parse_args()

    out = Path(args.out)
    img_dir = out / "images"
    img_dir.mkdir(parents=True, exist_ok=True)
    s = requests.Session()

    wb = openpyxl.load_workbook(args.excel)
    ws = wb.active
    header_row, cols = None, {}
    for row in ws.iter_rows(min_row=1, max_row=10):
        names = {str(c.value).strip(): c.column for c in row if c.value is not None}
        if "კოდი" in names and "მოდელის კოდი" in names:
            header_row, cols = row[0].row, names
            break
    if not header_row:
        sys.exit("სათაურის სტრიქონი ('კოდი', 'მოდელის კოდი') ვერ ვიპოვე")
    c_code, c_art = cols["კოდი"], cols["მოდელის კოდი"]
    c_model, c_color = cols.get("მოდელი"), cols.get("ფერი")
    c_status = ws.max_column + 1
    for i, title in enumerate(["სტატუსი", "საიტის ბმული", "სურათები"]):
        ws.cell(header_row, c_status + i, title).font = ws.cell(header_row, 1).font.copy(bold=True)

    print("Casco-ს კატალოგის ჩამოტვირთვა...")
    products = load_catalog(s, out / "casco_products.json")
    index = build_index(products)
    print(f"  სულ {len(products)} პროდუქტი, {len(index)} არტიკული")

    cache = {}  # არტიკული -> [jpeg bytes]
    stats = {k: 0 for k in FILLS}
    for r in range(header_row + 1, ws.max_row + 1):
        code = ws.cell(r, c_code).value
        art_full = ws.cell(r, c_art).value
        if code is None or art_full is None:
            continue  # კატეგორიის სტრიქონი
        code = str(code).strip()
        art = base_article(art_full)
        model = ws.cell(r, c_model).value if c_model else ""
        color = ws.cell(r, c_color).value if c_color else ""
        link, n_img = "", 0

        if args.odoo and (odoo_has(s, args.odoo, code) or odoo_has(s, args.odoo, art_full)):
            kind, status = "blue", "Odoo-ს საიტზე უკვე დევს — სურათები არ ჩამოიწერა"
        elif art in index:
            p, how = index[art][0]
            link = f"{SITE}/en/products/{p['handle']}"
            srcs = select_images(p, art)
            if not srcs:
                kind, status = "yellow", "საიტზე არის, მაგრამ სურათი არ აქვს"
            else:
                if art not in cache:
                    cache[art] = [fetch_jpeg(s, u, args.size) for u in srcs]
                for i, data in enumerate(cache[art]):
                    (img_dir / f"{code}{suffix(i)}.jpg").write_bytes(data)
                n_img = len(cache[art])
                kind = "green" if how == "sku" else "yellow"
                status = (f"ნაპოვნია: {p['title']}" if kind == "green"
                          else f"შესამოწმებელი: არტიკული მხოლოდ {'სურათის სახელში' if how == 'image' else 'ბმულში'} ემთხვევა — {p['title']}")
        else:
            cands = name_candidates(model, products)
            if cands:
                kind = "yellow"
                status = ("შესამოწმებელი: ეს მოდელი საიტზე არის, მაგრამ ამ არტიკულით (ფერით) არა. "
                          "მსგავსი: " + "; ".join(p["title"] for p in cands))
                link = f"{SITE}/en/products/{cands[0]['handle']}"
            else:
                kind, status = "red", "საიტზე ვერ ვიპოვე (ძველი მოდელი / აღარ იყიდება)"

        stats[kind] += 1
        for c in range(1, c_status + 3):
            ws.cell(r, c).fill = FILLS[kind]
        ws.cell(r, c_status, status)
        if link:
            ws.cell(r, c_status + 1, link).hyperlink = link
        ws.cell(r, c_status + 2, n_img or None)
        print(f"  [{kind:6}] {code} {art_full} {model} {color}: {status}")

    ws.column_dimensions[openpyxl.utils.get_column_letter(c_status)].width = 60
    ws.column_dimensions[openpyxl.utils.get_column_letter(c_status + 1)].width = 45
    wb.save(out / "Casco - შედეგი.xlsx")
    print("\nშედეგი:", stats, "| სურათები:", len(list(img_dir.glob('*.jpg'))))


if __name__ == "__main__":
    main()
