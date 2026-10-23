"""მომწოდებლების სურათების ჩამოწერა ოფიციალური საიტებიდან Excel-ის არტიკულების მიხედვით.

ბრენდი ავტომატურად ირჩევა Excel-ის სვეტიდან „ფირმა“ / „ბრენდი“ (ან --brand).
ახალი Shopify-ბრენდის დასამატებლად საკმარისია BRANDS-ში ერთი ჩანაწერი.

გამოყენება:
    python3 supplier_images.py FILE.xlsx --out OUT_DIR [--brand lenz] [--size 1500] [--odoo https://ODOO_DOMAIN]

შედეგი OUT_DIR-ში:
    images/<კოდი>.jpg, <კოდი>_A.jpg, ...      (პირველი: მთავარი სურათი)
    <ბრენდი>_result.xlsx                      (🟢 ნაპოვნია / 🟡 შესამოწმებელი / 🔴 ვერ ვიპოვე / 🔵 Odoo-ში უკვე არის)

Excel-ში უკვე მწვანედ მონიშნული სტრიქონები (სურათი უკვე აქვს) გამოიტოვება.
"""
import argparse
import io
import json
import re
import sys
import time
from pathlib import Path
from urllib.parse import quote

import openpyxl
import requests
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/141.0 Safari/537.36"}

FILLS = {
    "green": PatternFill("solid", fgColor="FFC6EFCE", bgColor="FFC6EFCE"),
    "yellow": PatternFill("solid", fgColor="FFFFEB9C", bgColor="FFFFEB9C"),
    "red": PatternFill("solid", fgColor="FFFFC7CE", bgColor="FFFFC7CE"),
    "blue": PatternFill("solid", fgColor="FFBDD7EE", bgColor="FFBDD7EE"),
}
ICONS = {"green": "🟢", "yellow": "🟡", "red": "🔴", "blue": "🔵", "skip": "⚪"}

# Excel-ის სვეტების შესაძლო სახელები
COLS = {
    "code": ["კოდი"],
    "article": ["მოდელის კოდი", "მომწ. კოდი", "მომწოდებლის კოდი"],
    "model": ["მოდელი", "სახელი"],
    "color": ["ფერი", "ფერი / ლინზა"],
    "colorcode": ["ფერის კოდი"],
    "note": ["შენიშვნა"],
    "upc": ["UPC", "UPC Code", "შტრიხკოდი"],
    "brand": ["ფირმა", "ბრენდი"],
}

# ინგლისური ფერი (Excel) -> გერმანული (საიტი)
COLOR_DE = {
    "BLACK": "schwarz", "WHITE": "weiss", "GREY": "grau", "GRAY": "grau",
    "LIGHTGREY": "hellgrau", "LIGHTGRAY": "hellgrau", "DARKGREY": "dunkelgrau",
    "BLUE": "blau", "NAVY": "marine", "RED": "rot", "PINK": "pink", "ROSE": "rosa",
    "GREEN": "gruen", "OLIVE": "oliv", "LIME": "lime", "YELLOW": "gelb",
    "ORANGE": "orange", "BROWN": "braun", "BEIGE": "beige", "ANTHRACITE": "anthrazit",
    "PURPLE": "violett", "VIOLET": "violett", "TURQUOISE": "tuerkis", "BERRY": "beere",
    "PETROL": "petrol", "SILVER": "silber",
}


# ---------- ბრენდები ----------

def casco_base(s):
    m = re.search(r"(\d{2})\.(\d{4})", str(s or "").replace(",", "."))
    return f"{m.group(1)}.{m.group(2)}" if m else None


class Casco:
    """casco-helmets.com: SKU 07.1902.L — '07.1902' = მოდელი+ფერი, ბოლო = ზომა."""
    name, site = "Casco", "https://casco-helmets.com"
    check_color = False  # ფერი არტიკულშივეა

    @staticmethod
    def article(excel_value):
        return casco_base(excel_value)

    @staticmethod
    def split_sku(sku, product=None):
        a = casco_base(sku)
        return (a, a) if a else (None, None)

    @staticmethod
    def image_key(fname, article):
        m = re.match(r"(\d{2}\.\d{4})_", fname)
        return m.group(1) if m else None

    @staticmethod
    def handle_article(handle):
        m = re.search(r"(\d{2})-(\d{4})(?:-|$)", handle)
        return f"{m.group(1)}.{m.group(2)}" if m else None


class Lenz:
    """lenzproducts.com: SKU 090_64 — '090' = მოდელი, '6' = ფერი, '4' = ზომა.
    სურათის სახელში ფერი ასე წერია: ...-090-60-1.webp"""
    name, site = "Lenz", "https://lenzproducts.com"
    check_color = True

    @staticmethod
    def norm(a):
        a = str(a or "").strip().upper()
        return a.lstrip("0") or a

    @classmethod
    def article(cls, excel_value):
        return cls.norm(excel_value) if excel_value not in (None, "") else None

    @classmethod
    def split_sku(cls, sku, product=None):
        if "_" not in (sku or ""):
            return None, None
        art, rest = sku.split("_", 1)
        return cls.norm(art), rest[:-1]

    @classmethod
    def image_key(cls, fname, article):
        m = re.search(rf"(?<![0-9a-z])0*{re.escape(article.lower())}[-_](\d+)0(?=[-_.])",
                      fname.lower())
        return m.group(1) if m else None

    @staticmethod
    def handle_article(handle):
        return None


class Shred:
    """shredoptics.com: Excel-ის „მომწ. კოდი“ = ვარიანტის სრული SKU (GOGRAO21A, HENOTP12L).
    სურათის სახელი ფერის კოდით იწყება: GOGRAO21A_21.jpg, HENOTP12_11.jpg (ზომის გარეშე)."""
    name, site = "Shred", "https://shredoptics.com"
    check_color = False  # SKU-ში ფერიცაა და ზომაც
    _prefixes = {}

    @staticmethod
    def article(excel_value):
        return str(excel_value).strip().upper() if excel_value not in (None, "") else None

    @classmethod
    def image_key(cls, fname, article=None):
        m = re.match(r"([A-Z]{2}[A-Z0-9]{4,})_", fname)
        return m.group(1) if m else None

    @classmethod
    def split_sku(cls, sku, product=None):
        sku = str(sku or "").strip().upper()
        if not sku:
            return None, None
        key = sku
        if product is not None:
            if product["id"] not in cls._prefixes:
                cls._prefixes[product["id"]] = {
                    k for k in (cls.image_key(image_name(i["src"])) for i in product.get("images", [])) if k}
            pref = [p for p in cls._prefixes[product["id"]] if sku.startswith(p)]
            if pref:
                key = max(pref, key=len)
        return sku, key

    @staticmethod
    def handle_article(handle):
        return None


class Schreuderssport:
    """schreuderssport.com (Shopify არ არის): პროდუქტის გვერდი /en/product/<არტიკული>.
    ფერები გვერდზე ასე წერია: 'Black/Red (ZWR)'; ფერის კოდი Excel-ის „შენიშვნა“-შიცაა.
    სურათები: <არტ>.jpg (ყველა ფერი ერთად), <არტ>_<ფერი>[_...].jpg, <არტ>_<აღწერა>.jpg (ზოგადი)."""
    name, site = "Schreuderssport", "https://www.schreuderssport.com"
    _pages = {}

    @classmethod
    def page(cls, session, art):
        if art not in cls._pages:
            html = get(f"{cls.site}/en/product/{quote(art)}", session).text
            cls._pages[art] = None if "<title>Product not found" in html else html
        return cls._pages[art]

    @staticmethod
    def color_names(name):
        out = []
        for part in str(name or "").lower().replace("gray", "grey").split("/"):
            k = re.sub(r"[^a-z]", "", part)
            if k:
                out.append(k)
        return sorted(out)

    @classmethod
    def resolve(cls, session, art_raw, note, color, model):
        art = str(art_raw).strip()
        link = f"{cls.site}/en/product/{quote(art)}"
        html = cls.page(session, art)
        if not html:
            return {"kind": "red", "status": "საიტზე ვერ ვიპოვე (ძველი მოდელი / აღარ იყიდება)"}
        title = re.search(r"<title>[^<]*? - ([^<]*?) - Design", html)
        title = title.group(1).strip() if title else art
        colors = {}
        for td in re.findall(r"<td>([^<]*\([A-Z0-9]{2,4}\)(?:<br />[^<]*\([A-Z0-9]{2,4}\))*)</td>", html):
            for item in td.split("<br />"):
                mm = re.match(r"\s*(.+?)\s*\(([A-Z0-9]{2,4})\)\s*$", item)
                if mm:
                    colors[mm.group(2)] = mm.group(1)
        m = re.search(r"imagedelivery\.net/([^/]+)/products/", html)
        names = []
        for n in re.findall(rf"products/({re.escape(art)}(?:_[^\"/ ]*)?\.jpg)", html, re.I):
            if n not in names:
                names.append(n)
        if not m or not names:
            return {"kind": "yellow", "status": f"საიტზე არის, მაგრამ სურათი არ აქვს — {title}", "link": link}

        note = str(note or "").strip().upper()
        want = cls.color_names(color)
        code, why = None, ""
        if note in colors:
            code = note
        else:
            exact = [c for c, nm in colors.items() if cls.color_names(nm) == want]
            if exact:
                code = exact[0]
            elif len(colors) == 1 and set(want) & set(cls.color_names(next(iter(colors.values())))):
                code = next(iter(colors))
                why = f" (Excel-ში '{color}', საიტზე '{colors[code]}')"
        if colors and not code:
            return {"kind": "yellow", "link": link,
                    "status": f"შესამოწმებელი: ფერი '{color}' ზუსტად ვერ დავადგინე. საიტზე არის: "
                              + "; ".join(f"{nm} ({c})" for c, nm in colors.items())}

        def parts(n):
            return n[len(art):-4].lstrip("_").split("_")

        if code:
            specific = [n for n in names if parts(n)[0].upper() == code]
            generic = [n for n in names if parts(n)[0] and parts(n)[0].upper() not in colors
                       and not re.fullmatch(r"[A-Z0-9]{3}", parts(n)[0])]
            bare = [n for n in names if not parts(n)[0]]
            chosen = specific + generic if specific else bare + generic
        else:
            chosen = names
        base = f"https://imagedelivery.net/{m.group(1)}/products/"
        return {"kind": "green", "link": link, "key": (art, code),
                "srcs": [base + n + "/public" for n in chosen],
                "status": f"ნაპოვნია: {title}" + (f" — {colors[code]} ({code})" if code else "") + why}


class Salomon:
    """salomon.com (Next.js + PerimeterX: გვერდები ბლოკავს). კატალოგს ვკითხულობთ საიტისავე Algolia-ს
    ძიებიდან: ყოველ პროდუქტს აქვს swatches {მომწ. კოდი (L47270100): image, imageBack, imagesExtra}.
    სურათები cdn.dam.salomon.com-ზეა (PNG 2000px+)."""
    name, site = "Salomon", "https://www.salomon.com"
    APP = "K1N6GB06IP"
    LOCALES = ["en-GB", "en-US", "de-DE", "fr-FR", "it-IT", "es-ES", "en-CA", "pl-PL", "sv-SE"]
    _key, _hits, _queried = None, [], set()

    @classmethod
    def _query(cls, session, loc, text, hits_per_page=1000, page=0):
        if cls._key is None:
            r = session.post(f"{cls.site}/api/algolia/secured-key/b", json={"subSegment": "D2C"},
                             headers={"User-Agent": "Mozilla/5.0"}, timeout=60)  # სრულ Chrome UA-ს PerimeterX ბლოკავს
            r.raise_for_status()
            cls._key = r.json()["apiKey"]
        r = session.post(f"https://{cls.APP}-dsn.algolia.net/1/indexes/prod_{loc}_products/query",
                         headers={"X-Algolia-Application-Id": cls.APP, "X-Algolia-API-Key": cls._key},
                         json={"query": text, "hitsPerPage": hits_per_page, "page": page}, timeout=60)
        return r.json() if r.ok else {}

    @classmethod
    def _search(cls, session, text):
        if text in cls._queried:
            return
        cls._queried.add(text)
        for loc in cls.LOCALES:
            res = cls._query(session, loc, text)
            cls._hits += [dict(h, _loc=loc) for h in res.get("hits", [])]
            if not text:  # ცარიელ ძიებაზე Algolia ბევრ გვერდს აბრუნებს
                for page in range(1, res.get("nbPages", 1)):
                    cls._hits += [dict(h, _loc=loc) for h in cls._query(session, loc, text, page=page).get("hits", [])]

    @staticmethod
    def model_tokens(name):
        """'X-ADVENTURE RECON GTX W' და 'X-ADVENTURE RECON GORE-TEX' -> ერთნაირი სიტყვები."""
        name = re.sub(r"\bGTX\b", "GORE TEX", str(name).upper())
        name = re.sub(r"\bWP\b", "WATERPROOF", name)
        return [t for t in tokens(name) if t not in ("W", "J", "M", "WIDE", "SET", "SHOES", "BAG")]

    @classmethod
    def _find(cls, mat):
        for h in cls._hits:
            sw = (h.get("swatches") or {}).get(mat)
            if sw:
                return h, sw
        return None, None

    @classmethod
    def resolve(cls, session, art_raw, note, color, model):
        mat = str(art_raw).strip().upper()
        cls._search(session, "")
        h, sw = cls._find(mat)
        words = [w for w in str(model or "").upper().split() if w not in ("SHOES", "W", "J", "M", "WIDE")]
        if not sw and words:
            for n in (3, 2):
                cls._search(session, " ".join(words[:n]))
            h, sw = cls._find(mat)
        if sw:
            base = h.get("imageBaseUrl", "https://cdn.dam.salomon.com")
            paths = [sw.get("image"), sw.get("imageBack")] + list(sw.get("imagesExtra") or [])
            srcs = []
            for pth in paths:
                if pth and mat in pth and re.search(r"\.(png|jpe?g|webp)(\?|$)", pth, re.I):
                    u = base + pth
                    if u not in srcs:
                        srcs.append(u)
            link = re.sub(r"/L[0-9A-Z]{8}$", "", h.get("url", "")) + f"/{mat}"
            if not srcs:
                return {"kind": "yellow", "link": link, "status": f"საიტზე არის, მაგრამ სურათი არ აქვს — {h.get('modelName')}"}
            return {"kind": "green", "link": link, "key": mat, "srcs": srcs,
                    "status": f"ნაპოვნია: {h.get('modelName')} ({h.get('detailedTitle', '')})"}
        cands = {}
        want = cls.model_tokens(" ".join(words))
        for x in cls._hits:
            have = cls.model_tokens(x.get("modelName", ""))
            if want and have == want:
                cands.setdefault(x["modelName"], set()).update((x.get("swatches") or {}).keys())
        if cands:
            return {"kind": "yellow",
                    "status": "შესამოწმებელი: ეს კოდი საიტზე არ არის, მაგრამ მოდელი სხვა ფერებით არის: "
                              + "; ".join(f"{k} ({', '.join(sorted(v))})" for k, v in cands.items())}
        return {"kind": "red", "status": "საიტზე ვერ ვიპოვე (Salomon-ის ონლაინ კატალოგში ეს მოდელი არ არის)"}


class Oakley:
    """oakley.com (Akamai ბლოკავს პროდუქტის გვერდებს). საიტის სწრაფი ძიება (/search/autocomplete)
    ღიაა და ვარიანტს UPC-ით პოულობს მთავარ სურათთან ერთად. დამატებითი ხედების მისამართებს
    UUID აქვს, ამიტომ აქედან მხოლოდ მთავარი სურათი მოდის."""
    name, site = "Oakley", "https://www.oakley.com"

    @classmethod
    def resolve(cls, session, art_raw, note, color, model, upc=None):
        upc = str(upc or "").strip()
        if not upc:
            return {"kind": "red", "status": "UPC არ არის მითითებული"}
        r = session.get(f"{cls.site}/en-us/search/autocomplete", params={"term": upc},
                        headers={"User-Agent": "Mozilla/5.0"}, timeout=60)
        prods = (r.json().get("products") or {}).values() if r.ok else []
        hit = next((x for x in prods if str(x.get("code")) == upc), None)
        if not hit or not hit.get("imagelarge"):
            return {"kind": "red", "status": "საიტზე ვერ ვიპოვე (ამ UPC-ით Oakley-ის საიტზე არაფერი იძებნება)"}
        img = re.sub(r"\?.*$", "", hit["imagelarge"]) + "?impolicy=PI20_Oakley_Resize&wid=2000"
        link = cls.site + "/en-us" + hit.get("url", "")
        kind, why = "green", ""
        m = re.search(r"__([0-9A-Z]{6})", img)
        if m and str(note or "").strip() and not str(note).strip().upper().startswith(m.group(1)):
            kind, why = "yellow", f" (შესამოწმებელი: სურათის ფერის კოდი {m.group(1)}, Excel-ში {note})"
        return {"kind": kind, "link": link, "key": upc, "srcs": [img],
                "status": f"ნაპოვნია: {hit.get('name')}{why} — მხოლოდ მთავარი სურათი"}


class Rehall:
    """rehall.com (Shopify). SKU-ები EAN-ია, მაგრამ სურათის სახელი = <მოდელის კოდი>_<ფერის კოდი>_<ნომერი>.jpg
    (60920_3015_01.jpg) — ზუსტად ინვოისის „Item Code“ და ფერის კოდი."""
    name, site = "Rehall", "https://rehall.com"
    _imgs, _names = None, {}

    @classmethod
    def _load(cls, session):
        if cls._imgs is not None:
            return
        cls._imgs, page = {}, 1
        while True:
            batch = get(f"{cls.site}/products.json?limit=250&page={page}", session).json()["products"]
            if not batch:
                break
            for p in batch:
                ci = color_index(p)
                vcolor = {v["id"]: v.get(f"option{ci + 1}") for v in p.get("variants", [])} if ci is not None else {}
                for i in p.get("images", []):
                    m = re.match(r"(\d{4,6})_(\d{3,5})_(\d{1,3})", image_name(i["src"]))
                    if m:
                        cls._imgs.setdefault((m.group(1), m.group(2)), {}).setdefault(int(m.group(3)), (i["src"], p))
                        for vid in i.get("variant_ids") or []:  # ფერის სახელი -> ფერის კოდი
                            if vcolor.get(vid):
                                cls._names[(m.group(1), vcolor[vid].strip().lower())] = m.group(2)
            page += 1
            time.sleep(0.5)

    @classmethod
    def resolve(cls, session, art_raw, note, color, model):
        cls._load(session)
        art, cc = str(art_raw).strip(), str(note or "").strip()
        found, why = cls._imgs.get((art, cc)), ""
        if not found and color:
            alt = cls._names.get((art, str(color).strip().lower()))
            if alt:
                found, why = cls._imgs.get((art, alt)), f" (შესამოწმებელი: ფერის კოდი Excel-ში {cc}, საიტზე {alt})"
        if not found:
            others = sorted(c for (a, c) in cls._imgs if a == art)
            if others:
                return {"kind": "yellow", "status": f"შესამოწმებელი: მოდელი საიტზე არის, მაგრამ ფერი {cc} არა. საიტზე: {', '.join(others)}"}
            return {"kind": "red", "status": "საიტზე ვერ ვიპოვე"}
        srcs = [found[k][0] for k in sorted(found)]
        p = found[min(found)][1]
        return {"kind": "yellow" if why else "green", "link": f"{cls.site}/products/{p['handle']}", "key": (art, cc),
                "srcs": srcs, "status": f"ნაპოვნია: {p['title']} — {color or cc}{why}"}


BRANDS = {"rehall": Rehall, "oakley": Oakley, "casco": Casco, "lenz": Lenz, "shred": Shred, "schreuderssport": Schreuderssport, "salomon": Salomon}


# ---------- ქსელი ----------

def get(url, session, **kw):
    for attempt in range(6):
        try:
            r = session.get(url, headers=UA, timeout=60, **kw)
        except requests.ConnectionError:
            if attempt == 5:
                raise
            time.sleep(2 ** attempt)
            continue
        if r.status_code in (429, 502, 503):
            time.sleep(2 ** attempt + 1)
            continue
        r.raise_for_status()
        return r
    r.raise_for_status()
    return r


def load_catalog(brand, session, cache):
    if cache.exists():
        return json.loads(cache.read_text(encoding="utf-8"))
    products, page = [], 1
    while True:
        batch = get(f"{brand.site}/products.json?limit=250&page={page}", session).json()["products"]
        if not batch:
            break
        products += batch
        print(f"  კატალოგი: გვერდი {page}, სულ {len(products)} პროდუქტი")
        page += 1
        time.sleep(0.5)
    cache.write_text(json.dumps(products, ensure_ascii=False), encoding="utf-8")
    return products


# ---------- შედარება ----------

def image_name(src):
    return src.split("?")[0].rsplit("/", 1)[-1]


def color_index(product):
    for i, o in enumerate(product.get("options", [])):
        if "farb" in o["name"].lower() or "color" in o["name"].lower():
            return i
    return None


def build_groups(brand, products):
    """არტიკული -> [{product, key, color, variant_ids}] (ერთი ჩანაწერი = პროდუქტის ერთი ფერი)."""
    index = {}
    for p in products:
        ci = color_index(p)
        groups = {}
        for v in p.get("variants", []):
            art, key = brand.split_sku(v.get("sku"), p)
            if not art:
                continue
            g = groups.setdefault((art, key), {
                "product": p, "article": art, "key": key, "variant_ids": set(),
                "color": v.get(f"option{ci + 1}") if ci is not None else "",
            })
            g["variant_ids"].add(v["id"])
        if not groups and brand.handle_article(p.get("handle", "")):
            art = brand.handle_article(p["handle"])
            groups[(art, art)] = {"product": p, "article": art, "key": art,
                                  "variant_ids": set(), "color": ""}
        for g in groups.values():
            index.setdefault(g["article"], []).append(g)
    return index


def images_for(brand, group):
    """ამ ფერის სურათები. ფერი დგინდება: ვარიანტთან მიბმით, სახელში კოდით,
    ან (თუ სახელში კოდი არ წერია) წინა ფერიანი სურათის მიხედვით."""
    p = group["product"]
    imgs = sorted(p.get("images", []), key=lambda i: i.get("position", 0))
    keys_of = {}
    for v in p.get("variants", []):
        art, key = brand.split_sku(v.get("sku"), p)
        if art:
            keys_of[v["id"]] = key
    all_keys = set(keys_of.values())
    if len(all_keys) <= 1:
        return [i["src"] for i in imgs]
    out, current = [], None
    for i in imgs:
        linked = {keys_of[x] for x in i.get("variant_ids") or [] if x in keys_of}
        named = brand.image_key(image_name(i["src"]), group["article"])
        if linked:
            current = group["key"] if group["key"] in linked else next(iter(linked))
            mine = group["key"] in linked
        elif named is not None:
            if named not in all_keys:
                continue  # სხვა პროდუქტის ფერი
            current = named
            mine = named == group["key"]
        else:
            mine = current == group["key"]
        if mine and i["src"] not in out:
            out.append(i["src"])
    return out


def norm_de(s):
    s = str(s).lower().replace("ß", "ss").replace("ä", "ae").replace("ö", "oe").replace("ü", "ue")
    return s


def color_tokens_excel(color):
    out = []
    for part in str(color or "").upper().split("/"):
        k = re.sub(r"[^A-Z]", "", part)
        if k:
            out.append(COLOR_DE.get(k, k.lower()))
    return out


def color_tokens_site(color):
    return [norm_de(t).strip() for t in str(color or "").split("/") if t.strip()]


def color_score(excel_color, site_color):
    a, b = color_tokens_excel(excel_color), color_tokens_site(site_color)
    if not a or not b:
        return 0.0
    if sorted(a) == sorted(b):
        return 1.0
    common = len(set(a) & set(b))
    return common / max(len(set(a)), len(set(b)))


def tokens(s):
    s = re.sub(r"[^A-Z0-9]+", " ", str(s).upper().replace("+", " PLUS "))
    return [re.sub(r"(.)\1+", r"\1", t) for t in s.split()
            if t not in {"ML", "MAGNET", "LINK", "CASCO", "LENZ"}]


def title_score(model, title):
    want, have = set(tokens(model)), tokens(title)
    return sum(t in want for t in have) - 2 * sum(t not in want for t in have)


def name_candidates(model, products, limit=3):
    want = tokens(model)
    found = []
    for p in products:
        title = tokens(p.get("title", ""))
        have = set(title + tokens(p.get("handle", "")))
        if want and title[:1] == want[:1] and all(t in have for t in want):
            found.append(p)
    return found[:limit]


def suffix(i):
    """0 -> '', 1 -> '_A', 26 -> '_Z', 27 -> '_AA'"""
    s = ""
    while i > 0:
        i -= 1
        s = chr(65 + i % 26) + s
        i //= 26
    return "_" + s if s else ""


def is_marked_green(row):
    for c in row:
        f = c.fill
        if f is not None and f.fill_type == "solid" and f.fgColor is not None and f.fgColor.type == "rgb":
            rgb = str(f.fgColor.rgb).upper()[-6:]
            r, g, b = int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16)
            if g > 120 and g > r + 40 and g > b + 40:
                return True
    return False


# ---------- Odoo-ს საჯარო საიტი ----------

def odoo_has(session, odoo, code):
    r = get(f"{odoo}/shop?search={quote(str(code))}", session)
    if "No results" in r.text:
        return False
    return bool(re.search(r'href="(?:/[a-z]{2}(?:_[A-Z]{2})?)?/shop/(?!cart|checkout|category|page|wishlist)[^"?#]+-\d+',
                          r.text))


# ---------- სურათები ----------

def fetch_jpeg(session, src, size):
    url = "https:" + src if src.startswith("//") else src
    if size and ("cdn.shopify.com" in url or "/cdn/shop/" in url):
        url += ("&" if "?" in url else "?") + f"width={size}"
    im = Image.open(io.BytesIO(get(url, session).content))
    im.load()
    if size and max(im.size) > size:
        im.thumbnail((size, size), Image.LANCZOS)
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.split()[-1])
        im = bg
    elif im.mode != "RGB":
        im = im.convert("RGB")
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=92, optimize=True)
    return buf.getvalue()


# ---------- მთავარი ----------

def find_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=15):
        names = {str(c.value).strip(): c.column for c in row if c.value is not None}
        cols = {k: next((names[n] for n in v if n in names), None) for k, v in COLS.items()}
        if cols["article"]:
            return row[0].row, cols
    sys.exit("სათაურის სტრიქონი ვერ ვიპოვე (საჭიროა 'კოდი' და 'მოდელის კოდი' / 'მომწ. კოდი')")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--out", required=True)
    ap.add_argument("--brand", choices=sorted(BRANDS))
    ap.add_argument("--size", type=int, default=1500, help="მაქს. ზომა პიქსელებში; 0 = ორიგინალი")
    ap.add_argument("--odoo", default=None)
    ap.add_argument("--sheet", default=None, help="ფურცლის სახელი (ნაგულისხმევად პირველი)")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.excel)
    ws = wb[args.sheet] if args.sheet else wb.active
    header_row, cols = find_header(ws)

    brand_key = args.brand
    if not brand_key and cols["brand"]:
        for r in range(header_row + 1, ws.max_row + 1):
            v = str(ws.cell(r, cols["brand"]).value or "").strip().lower()
            if v in BRANDS:
                brand_key = v
                break
    if not brand_key:
        sys.exit(f"ბრენდი ვერ დავადგინე; მიუთითეთ --brand ({', '.join(sorted(BRANDS))})")
    brand = BRANDS[brand_key]

    out = Path(args.out)
    img_dir = out / "images"
    img_dir.mkdir(parents=True, exist_ok=True)
    s = requests.Session()

    c_status = ws.max_column + 1
    bold = ws.cell(header_row, 1).font.copy(bold=True)
    for i, title in enumerate(["სტატუსი", "საიტის ბმული", "სურათები"]):
        ws.cell(header_row, c_status + i, title).font = bold

    products, index = [], {}
    if not hasattr(brand, "resolve"):
        print(f"{brand.name}: კატალოგის ჩამოტვირთვა...")
        products = load_catalog(brand, s, out / f"{brand_key}_products.json")
        index = build_groups(brand, products)
        print(f"  სულ {len(products)} პროდუქტი, {len(index)} არტიკული")

    cache = {}
    stats = {k: 0 for k in ICONS}
    for r in range(header_row + 1, ws.max_row + 1):
        row = ws[r]
        art_raw = ws.cell(r, cols["article"]).value
        code = ws.cell(r, cols["code"]).value if cols["code"] else None
        if code in (None, ""):
            # Odoo-ს კოდი თუ არ არის: UPC (ერთი ფერი), თორემ მომწოდებლის კოდი
            upc = ws.cell(r, cols["upc"]).value if cols["upc"] else None
            cc = ws.cell(r, cols["colorcode"]).value if cols["colorcode"] else None
            code = upc if upc not in (None, "") else f"{art_raw}_{cc}" if cc not in (None, "") else art_raw
        if code in (None, "") or art_raw in (None, ""):
            continue
        code = str(code).strip()
        model = ws.cell(r, cols["model"]).value if cols["model"] else ""
        color = ws.cell(r, cols["color"]).value if cols["color"] else ""
        link, n_img, kind = "", 0, None

        if is_marked_green(row):
            kind, status = "skip", "უკვე აქვს სურათი (მწვანედ იყო მონიშნული) — გამოტოვდა"
        elif args.odoo and odoo_has(s, args.odoo, code):
            kind, status = "blue", "Odoo-ს საიტზე უკვე დევს — სურათები არ ჩამოიწერა"
        elif hasattr(brand, "resolve"):
            note = ws.cell(r, cols["note"]).value if cols["note"] else ""
            extra = {}
            if brand is Oakley:
                extra["upc"] = ws.cell(r, cols["upc"]).value if cols["upc"] else None
                note = ws.cell(r, cols["colorcode"]).value if cols["colorcode"] else ""  # 7064H200
            if brand is Rehall:
                note = ws.cell(r, cols["colorcode"]).value if cols["colorcode"] else ""  # 3015
            res = brand.resolve(s, art_raw, note, color, model, **extra)
            kind, status, link = res["kind"], res["status"], res.get("link", "")
            if res.get("srcs"):
                if res["key"] not in cache:
                    cache[res["key"]] = [fetch_jpeg(s, u, args.size) for u in res["srcs"]]
                for i, data in enumerate(cache[res["key"]]):
                    (img_dir / f"{code}{suffix(i)}.jpg").write_bytes(data)
                n_img = len(cache[res["key"]])
        else:
            art = brand.article(art_raw)
            groups = index.get(art, [])
            chosen, sure = None, True
            if groups and brand.check_color and color:
                scored = sorted(((color_score(color, g["color"]), g) for g in groups),
                                key=lambda t: -t[0])
                best, chosen = scored[0]
                if best < 1.0:
                    sure = False
            elif groups:
                # ერთი კოდი რამდენიმე პროდუქტშია (მაგ. GRATIFY და GRATIFY+): ვირჩევთ სახელით
                chosen = max(groups, key=lambda g: title_score(model, g["product"]["title"]))

            if chosen and sure:
                p = chosen["product"]
                link = f"{brand.site}/products/{p['handle']}"
                srcs = images_for(brand, chosen)
                if not srcs:
                    kind, status = "yellow", f"საიტზე არის, მაგრამ სურათი არ აქვს — {p['title']}"
                else:
                    ck = (p["id"], chosen["key"])
                    if ck not in cache:
                        cache[ck] = [fetch_jpeg(s, u, args.size) for u in srcs]
                    for i, data in enumerate(cache[ck]):
                        (img_dir / f"{code}{suffix(i)}.jpg").write_bytes(data)
                    n_img = len(cache[ck])
                    kind = "green"
                    status = f"ნაპოვნია: {p['title']}" + (f" — {chosen['color']}" if chosen["color"] else "")
            elif chosen:
                p = chosen["product"]
                link = f"{brand.site}/products/{p['handle']}"
                kind = "yellow"
                status = (f"შესამოწმებელი: მოდელი საიტზე არის, მაგრამ ფერი '{color}' ზუსტად არ ემთხვევა. "
                          "საიტზე არის: " + "; ".join(g["color"] for g in groups))
            else:
                cands = name_candidates(model, products)
                if cands:
                    kind = "yellow"
                    status = ("შესამოწმებელი: ეს მოდელი საიტზე არის, მაგრამ ამ არტიკულით არა. "
                              "მსგავსი: " + "; ".join(p["title"] for p in cands))
                    link = f"{brand.site}/products/{cands[0]['handle']}"
                else:
                    kind, status = "red", "საიტზე ვერ ვიპოვე (ძველი მოდელი / აღარ იყიდება)"

        stats[kind] += 1
        if kind in FILLS:
            for c in range(1, c_status + 3):
                ws.cell(r, c).fill = FILLS[kind]
        ws.cell(r, c_status, f"{ICONS[kind]} {status}")
        if link:
            ws.cell(r, c_status + 1, link).hyperlink = link
        ws.cell(r, c_status + 2, n_img or None)
        print(f"  [{kind:6}] {code} {art_raw} {model} {color}: {status}")

    ws.column_dimensions[get_column_letter(c_status)].width = 60
    ws.column_dimensions[get_column_letter(c_status + 1)].width = 45
    wb.save(out / f"{brand.name}_result.xlsx")
    print("\nშედეგი:", stats, "| სურათები:", len(list(img_dir.glob("*.jpg"))))


if __name__ == "__main__":
    main()
